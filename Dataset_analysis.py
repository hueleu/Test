#!/usr/bin/env python3
"""
vad_dataset_analysis.py
Phân tích chất lượng dataset VAD + mức độ nghiêm trọng lỗi theo ngữ cảnh.

Cấu trúc thư mục kỳ vọng:
    <model_root>/<Category>_<idx>.txt                        (raw YAMNet HOẶC đã rebin 0.5s)
    <gt_root>/<Category>/groundtruth/Gt_<category>_<idx>.txt (start end label)
GT label chấp nhận: speech / non-speech (hoặc 1 / 0). Phân cách: space, tab, dấu phẩy.

Model input (--model_mode, mặc định auto = tự nhận):
    raw      : output YAMNet (window 0.96s, hop 0.08s) → script tự rescore, pad, merge, rebin 0.5s
    rebinned : đã rebin 0.5s → đọc thẳng, KHÔNG post-process. Các dạng dòng chấp nhận:
               [index] start end score | [index] start end label | [index] start end score label
               (label = speech/non-speech hoặc 0/1; nếu chỉ có score thì dùng ngưỡng để ra nhãn)

Chạy:
    python vad_dataset_analysis.py --model_root Model --gt_root Groundtruth --out vad_report
"""
import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd

# ----------------------------------------------------------------------------- config
CFG = dict(
    threshold=0.6026,      # ngưỡng Youden
    hop=0.08,              # hop YAMNet (s)
    pad_before=0.10, pad_after=0.12,
    merge_gap=0.50,        # gộp segment có gap < 0.5s
    bin=0.5,               # độ phân giải GT (s)
    intra_gap_max=1.0,     # non-speech kẹp giữa speech, <= 1s  -> khoảng ngắt tự nhiên
    turn_gap_max=2.0,      # non-speech > 2s -> non-speech dài (vùng kiểm tra loại nhiễu)
    short_speech_max=1.0,  # speech <= 1s -> speech ngắn
    collar_bins=1,         # số bin quanh điểm chuyển coi là lỗi biên
    conf_hi=0.9, conf_lo=0.1,
    min_events=30,         # số đoạn tối thiểu / nhóm để kết luận có ý nghĩa
    n_boot=1000, seed=0,
)
# trọng số độ nghiêm trọng theo ngữ cảnh đoạn GT chứa lỗi
CONTEXT_W = {"ns_intra_gap": 0.1, "ns_edge": 0.5, "ns_turn_gap": 0.5, "ns_long": 1.0,
             "sp_short": 1.0, "sp_normal": 1.0}
BOUNDARY_FACTOR = 0.3      # lỗi nằm trong collar được nhân thêm hệ số này

SPEECH = {"speech", "1", "sp", "s"}
NONSPEECH = {"non-speech", "nonspeech", "non_speech", "0", "ns", "silence", "noise"}


# ----------------------------------------------------------------------------- IO
def _tok(line):
    return [t for t in re.split(r"[,\t; ]+", line.strip()) if t]


def _isnum(t):
    try:
        float(t)
        return True
    except ValueError:
        return False


def read_model(path):
    rows = []
    for line in Path(path).read_text(encoding="utf-8", errors="ignore").splitlines():
        nums = [float(t) for t in _tok(line) if _isnum(t)]
        if len(nums) >= 4:
            rows.append(nums[1:4])          # index, start, end, score
        elif len(nums) == 3:
            rows.append(nums)               # start, end, score
    return np.array(rows, dtype=float).reshape(-1, 3)


def read_gt(path):
    segs = []
    for line in Path(path).read_text(encoding="utf-8", errors="ignore").splitlines():
        tok = _tok(line)
        if not tok:
            continue
        words = [t.lower() for t in tok if not _isnum(t)]
        nums = [float(t) for t in tok if _isnum(t)]
        if words:
            w = words[-1]
            lab = 1 if w in SPEECH else 0 if w in NONSPEECH else None
            se = nums[-2:]
        else:
            lab = int(nums[-1]) if len(nums) >= 3 else None
            se = nums[-3:-1]
        if lab is None or len(se) < 2:
            continue
        segs.append((se[0], se[1], lab))
    return segs


def match_files(model_root, gt_root):
    pairs = []
    for mf in sorted(Path(model_root).glob("*.txt")):
        m = re.match(r"^(.*)_(\d+)$", mf.stem)
        if not m:
            continue
        cat, idx = m.group(1), m.group(2)
        cand = Path(gt_root) / cat / "groundtruth" / f"Gt_{cat.lower()}_{idx}.txt"
        if not cand.exists():
            target = f"gt_{cat.lower()}_{idx}.txt"
            hits = [p for p in Path(gt_root).rglob("*.txt") if p.name.lower() == target]
            cand = hits[0] if hits else None
        if cand is None:
            print(f"[WARN] không tìm thấy GT cho {mf.name}")
            continue
        pairs.append((cat, mf.stem, mf, cand))
    return pairs


# ----------------------------------------------------------------------------- helpers
def runs(a):
    a = np.asarray(a)
    if len(a) == 0:
        return []
    cut = np.flatnonzero(a[1:] != a[:-1]) + 1
    st, en = np.r_[0, cut], np.r_[cut, len(a)]
    return [(int(s), int(e), a[s]) for s, e in zip(st, en)]


def gt_to_bins(segs, b):
    T = max(e for _, e, _ in segs)
    nb = int(np.ceil(T / b - 1e-9))
    lab = np.full(nb, -1)
    c = (np.arange(nb) + 0.5) * b
    for s, e, l in segs:
        lab[(c >= s) & (c < e)] = l
    return lab


def read_frames(path):
    """Đọc file dạng frame/đoạn: [index] start end [score] [label]. Tự nhận cột theo nội dung."""
    rows, labs = [], []
    for line in Path(path).read_text(encoding="utf-8", errors="ignore").splitlines():
        tok = _tok(line)
        if not tok:
            continue
        words = [t.lower() for t in tok if not _isnum(t)]
        nums = [float(t) for t in tok if _isnum(t)]
        lab = np.nan
        if words:
            w = words[-1]
            if w in SPEECH:
                lab = 1.0
            elif w in NONSPEECH:
                lab = 0.0
            else:
                continue                                   # dòng tiêu đề
        if len(nums) < 2:
            continue
        rows.append(nums); labs.append(lab)
    if not rows:
        return pd.DataFrame(columns=["start", "end", "score", "label"], dtype=float)
    w = min(len(r) for r in rows)
    M = np.array([r[:w] for r in rows])
    labs = np.array(labs, dtype=float)

    def is_start_end(i):
        # end >= start mọi dòng, > start ở hầu hết dòng (cho phép frame 0s do làm tròn), start tăng dần
        if i + 1 >= w:
            return False
        d = M[:, i + 1] - M[:, i]
        return bool((d >= -1e-9).all() and (d > 1e-9).mean() >= 0.9
                    and (np.diff(M[:, i]) >= -1e-9).all())

    si = 1 if is_start_end(1) else 0 if is_start_end(0) else None
    if si is None:
        raise SystemExit(f"Không nhận ra cột start/end trong {path}")
    keep = M[:, si + 1] - M[:, si] > 1e-9                 # bỏ frame dài 0s
    M, labs = M[keep], labs[keep]
    extra = M[:, si + 2:]
    score = extra[:, 0].copy() if extra.shape[1] >= 1 else np.full(len(M), np.nan)
    if np.isnan(labs).all() and extra.shape[1] >= 2:
        labs = extra[:, 1].copy()                          # start end score label(0/1)
    if np.isnan(labs).all() and np.isfinite(score).all() and np.isin(score, [0.0, 1.0]).all():
        labs, score = score, np.full(len(M), np.nan)      # cột duy nhất là nhãn 0/1
    return pd.DataFrame(dict(start=M[:, si], end=M[:, si + 1], score=score, label=labs))


def is_frame_grid(df, b, tol=0.02):
    """True nếu mọi dòng (trừ dòng cuối có thể ngắn hơn) là 1 frame dài đúng b giây."""
    if len(df) == 0:
        return False
    dur = (df.end - df.start).values
    core = dur[:-1] if len(dur) > 1 else dur
    return bool(np.all(np.abs(core - b) < tol) and dur[-1] <= b + tol)


def frames_to_grid(df, col, b, n=None, who=""):
    """Đặt giá trị df[col] lên lưới bin b giây.
    Frame đúng b giây → đặt thẳng theo chỉ số (không rebin). Đoạn dài hơn → gán theo tâm bin."""
    if n is None:
        n = int(np.ceil(df.end.max() / b - 1e-9))
    out = np.full(n, np.nan)
    vals = df[col].values.astype(float)
    if is_frame_grid(df, b):
        k = np.round(df.start.values / b).astype(int)
        off = np.abs(df.start.values - k * b).max()
        if off > 0.01:
            print(f"[WARN] {who}: frame lệch lưới {b}s tối đa {off:.3f}s — kiểm tra mốc thời gian")
        if len(np.unique(k)) < len(k):
            print(f"[WARN] {who}: có frame trùng chỉ số, giữ frame xuất hiện sau")
        m = (k >= 0) & (k < n)
        out[k[m]] = vals[m]
    else:
        c = (np.arange(n) + 0.5) * b
        for s, e, v in zip(df.start.values, df.end.values, vals):
            out[(c >= s) & (c < e)] = v
    return out


def model_to_bins(win, nb, cfg):
    hop, b, thr = cfg["hop"], cfg["bin"], cfg["threshold"]
    T = max(win[:, 1].max(), nb * b)
    nh = int(np.ceil(T / hop))
    acc, cnt = np.zeros(nh), np.zeros(nh)
    for s, e, sc in win:                       # rescore từng hop = mean các window phủ hop
        i0, i1 = int(round(s / hop)), max(int(round(e / hop)), int(round(s / hop)) + 1)
        acc[i0:i1] += sc
        cnt[i0:i1] += 1
    hs = np.where(cnt > 0, acc / np.maximum(cnt, 1), 0.0)
    speech = hs >= thr
    segs = [[s * hop - cfg["pad_before"], e * hop + cfg["pad_after"]]
            for s, e, v in runs(speech) if v]
    merged = []
    for s, e in segs:
        s, e = max(s, 0.0), min(e, T)
        if merged and s - merged[-1][1] < cfg["merge_gap"]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    pred = np.zeros(nb, dtype=int)
    for i in range(nb):
        a0, a1 = i * b, (i + 1) * b
        ov = sum(max(0.0, min(a1, e) - max(a0, s)) for s, e in merged)
        pred[i] = int(ov >= 0.5 * b)
    hc = (np.arange(nh) + 0.5) * hop
    score = np.array([hs[(hc >= i * b) & (hc < (i + 1) * b)].mean()
                      if ((hc >= i * b) & (hc < (i + 1) * b)).any() else np.nan
                      for i in range(nb)])
    pred_raw = (np.nan_to_num(score) >= thr).astype(int)
    return pred, pred_raw, score


def classify_segments(gt, cfg):
    b = cfg["bin"]
    rs = runs(gt)
    seg_id = np.full(len(gt), -1)
    seg_type = np.array(["unlabeled"] * len(gt), dtype=object)
    info = []
    for j, (s, e, v) in enumerate(rs):
        dur = (e - s) * b
        if v == 1:
            t = "sp_short" if dur <= cfg["short_speech_max"] else "sp_normal"
        elif v == 0:
            prev_sp = j > 0 and rs[j - 1][2] == 1
            next_sp = j + 1 < len(rs) and rs[j + 1][2] == 1
            if dur > cfg["turn_gap_max"]:
                t = "ns_long"
            elif prev_sp and next_sp:
                t = "ns_intra_gap" if dur <= cfg["intra_gap_max"] else "ns_turn_gap"
            else:
                t = "ns_edge"
        else:
            t = "unlabeled"
        seg_id[s:e], seg_type[s:e] = j, t
        info.append(dict(seg=j, b0=s, b1=e, label=int(v), type=t, dur=dur))
    return seg_id, seg_type, info


def transition_dist(gt):
    n = len(gt)
    k = np.array([i for i in range(1, n) if gt[i] != gt[i - 1] and gt[i] >= 0 and gt[i - 1] >= 0])
    if len(k) == 0:
        return np.full(n, 10 ** 6)
    i = np.arange(n)[:, None]
    d = np.where(i >= k[None, :], i - k[None, :], k[None, :] - 1 - i)
    return d.min(axis=1)


def metrics(gt, pr):
    gt, pr = np.asarray(gt), np.asarray(pr)
    tp = int(((gt == 1) & (pr == 1)).sum()); fp = int(((gt == 0) & (pr == 1)).sum())
    fn = int(((gt == 1) & (pr == 0)).sum()); tn = int(((gt == 0) & (pr == 0)).sum())
    p = tp / (tp + fp) if tp + fp else np.nan
    r = tp / (tp + fn) if tp + fn else np.nan
    f1 = 2 * p * r / (p + r) if p == p and r == r and p + r else np.nan
    spec = tn / (tn + fp) if tn + fp else np.nan
    acc = (tp + tn) / max(tp + fp + fn + tn, 1)
    return dict(TP=tp, FP=fp, FN=fn, TN=tn, precision=p, recall=r, F1=f1,
                specificity=spec, accuracy=acc)


# ----------------------------------------------------------------------------- per file
def analyze_file(cat, name, mf, gf, cfg):
    b = cfg["bin"]
    gdf = read_frames(gf)
    if gdf.label.isna().all():
        raise SystemExit(f"GT {gf}: không tìm thấy cột nhãn speech/non-speech")
    g = frames_to_grid(gdf, "label", b, who=f"GT {name}")
    gt = np.where(np.isnan(g), -1, g).astype(int)
    if cfg["mode"] == "rebinned":
        mdf = read_frames(mf)
        if not is_frame_grid(mdf, b):
            print(f"[WARN] Model {name}: dòng không phải frame {b}s — gán theo tâm bin")
        score = frames_to_grid(mdf, "score", b, len(gt), f"Model {name}")
        if mdf.label.notna().any():
            p = frames_to_grid(mdf, "label", b, len(gt), f"Model {name}")
        else:
            p = np.where(np.isnan(score), np.nan, (score >= cfg["threshold"]).astype(float))
        pred = np.where(np.isnan(p), -1, p).astype(int)
        pred_raw = pred.copy()                            # không có bản raw ở chế độ này
        miss = int(((pred < 0) & (gt >= 0)).sum())
        if miss:
            print(f"[WARN] {name}: {miss} frame GT không có frame model tương ứng → bỏ qua")
    else:
        pred, pred_raw, score = model_to_bins(read_model(mf), len(gt), cfg)
    seg_id, seg_type, seg_info = classify_segments(gt, cfg)
    dist = transition_dist(gt)
    sp_idx = np.flatnonzero(gt == 1)
    fill = np.zeros_like(gt)
    if len(sp_idx):
        fill[sp_idx[0]:sp_idx[-1] + 1] = 1
    bins = pd.DataFrame(dict(category=cat, file=name, bin=np.arange(len(gt)),
                             t0=np.arange(len(gt)) * cfg["bin"], y=gt, pred=pred,
                             pred_raw=pred_raw, fill_all=fill, all_speech=1, score=score,
                             seg=seg_id, seg_type=seg_type, dist=dist))
    bins = bins[(bins.y >= 0) & (bins.pred >= 0)].copy()
    bins["err"] = (bins.pred != bins.y).astype(int)
    bins["in_collar"] = bins.dist < cfg["collar_bins"]
    bins["weight"] = bins.seg_type.map(CONTEXT_W).fillna(1.0) * np.where(bins.in_collar, BOUNDARY_FACTOR, 1.0)
    bins["w_err"] = bins.err * bins.weight

    segs = pd.DataFrame(seg_info)
    segs = segs[segs.label >= 0].copy()
    segs["category"], segs["file"] = cat, name
    cov = bins.groupby("seg").pred.mean()
    segs["pred_speech_frac"] = segs.seg.map(cov)

    # error runs + severity level
    key = np.where(bins.err.values == 1, bins.y.values + 1, 0)
    run_rows = []
    for s, e, v in runs(key):
        if v == 0:
            continue
        sub = bins.iloc[s:e]
        kind = "FP" if v == 1 else "FN"
        stype = sub.seg_type.mode().iloc[0]
        seg_row = segs[segs.seg == sub.seg.iloc[0]].iloc[0]
        if kind == "FN" and seg_row.pred_speech_frac == 0:
            lvl = "L4_missed_whole_speech"
        elif kind == "FP" and stype == "ns_long" and seg_row.pred_speech_frac == 1:
            lvl = "L4_whole_long_ns_as_speech"
        elif kind == "FP" and (sub.seg_type == "ns_intra_gap").all():
            lvl = "L0_gap_fill"
        elif sub.in_collar.all():
            lvl = "L1_boundary"
        elif e - s <= 2:
            lvl = "L2_short"
        else:
            lvl = "L3_segment"
        sc = sub.score.mean()
        conf = (kind == "FP" and sc >= cfg["conf_hi"]) or (kind == "FN" and sc <= cfg["conf_lo"])
        run_rows.append(dict(category=cat, file=name, t_start=sub.t0.iloc[0],
                             t_end=sub.t0.iloc[-1] + cfg["bin"], n_bins=e - s, kind=kind,
                             seg_type=stype, level=lvl, mean_score=sc, confident=conf,
                             w_err=sub.w_err.sum()))
    return bins, segs, pd.DataFrame(run_rows)


# ----------------------------------------------------------------------------- report
def fmt(df, **kw):
    return df.to_string(float_format=lambda x: f"{x:.3f}", **kw)


def boot_ci(bins, cfg):
    rng = np.random.default_rng(cfg["seed"])
    g = bins.assign(tp=(bins.y == 1) & (bins.pred == 1), fp=(bins.y == 0) & (bins.pred == 1),
                    fn=(bins.y == 1) & (bins.pred == 0),
                    nl=bins.seg_type == "ns_long",
                    fpl=(bins.seg_type == "ns_long") & (bins.pred == 1))
    per = g.groupby("file")[["tp", "fp", "fn", "nl", "fpl"]].sum().values
    if len(per) < 2:
        return (np.nan, np.nan), (np.nan, np.nan)
    f1s, fprs = [], []
    for _ in range(cfg["n_boot"]):
        s = per[rng.integers(0, len(per), len(per))].sum(axis=0)
        tp, fp, fn, nl, fpl = s
        f1s.append(2 * tp / max(2 * tp + fp + fn, 1))
        fprs.append(fpl / nl if nl else np.nan)
    q = lambda a: tuple(np.nanpercentile(a, [2.5, 97.5])) if np.isfinite(a).any() else (np.nan, np.nan)
    return q(np.array(f1s)), q(np.array(fprs, dtype=float))


def build_report(bins, segs, errs, cfg, ci):
    L, v = [], []
    add = L.append
    b = cfg["bin"]

    add("=" * 78); add("BÁO CÁO PHÂN TÍCH DATASET & MỨC ĐỘ LỖI VAD"); add("=" * 78)
    add(f"Files: {bins.file.nunique()} | Categories: {bins.category.nunique()} | "
        f"Tổng thời lượng: {len(bins) * b / 60:.1f} phút | Tỷ lệ speech: {bins.y.mean():.1%}")
    if cfg["has_raw"]:
        add(f"Model input: raw | Ngưỡng={cfg['threshold']}, pad={cfg['pad_before']}/{cfg['pad_after']}s, "
            f"merge<{cfg['merge_gap']}s, bin={b}s, collar={cfg['collar_bins']} bin")
    else:
        add(f"Model input: đã rebin {b}s (không post-process lại) | collar={cfg['collar_bins']} bin | "
            f"ngưỡng {cfg['threshold']} chỉ dùng cho file không có cột nhãn")

    # 1. structure
    add("\n[1] CẤU TRÚC ĐOẠN TRONG GROUND TRUTH")
    st = segs.groupby("type").agg(n_events=("seg", "size"), total_s=("dur", "sum"),
                                  mean_s=("dur", "mean")).reindex(list(CONTEXT_W)).dropna(how="all")
    ns_total = st.loc[st.index.str.startswith("ns"), "total_s"].sum()
    st["share_of_ns"] = np.where(st.index.str.startswith("ns"), st.total_s / ns_total, np.nan)
    st["đủ_mẫu"] = np.where(st.n_events >= cfg["min_events"], "OK", "THIẾU")
    add(fmt(st))
    add("\nSố đoạn theo category:")
    add(fmt(segs.pivot_table(index="category", columns="type", values="seg",
                             aggfunc="size", fill_value=0)))

    # 2. baselines
    add("\n[2] SO VỚI BASELINE NGÂY THƠ (toàn bộ bin)")
    names = {"pred": "model (post-proc)" if cfg["has_raw"] else "model (đã rebin)",
             "pred_raw": "model (raw, không post-proc)",
             "fill_all": "fill-all (onset đầu→offset cuối)", "all_speech": "all-speech"}
    keys = [k for k in names if cfg["has_raw"] or k != "pred_raw"]
    bl = pd.DataFrame({k: metrics(bins.y, bins[k]) for k in keys}).T
    gap = bl.loc["pred"].F1 - bl.loc["fill_all"].F1
    bl.index = [names[k] for k in keys]
    add(fmt(bl[["precision", "recall", "F1", "specificity", "accuracy"]]))
    add(f"→ Chênh F1 model − fill-all = {gap:+.3f}")
    if gap < 0.05:
        v.append(f"Dataset DỄ: fill-all chỉ kém model {gap:.3f} F1, metric toàn cục không phân biệt được model tốt/tầm thường.")

    # 3. stratified
    add("\n[3] TỶ LỆ LỖI PHÂN TẦNG THEO NGỮ CẢNH ĐOẠN")
    rows = []
    for t, g in bins.groupby("seg_type"):
        if t == "unlabeled":
            continue
        rows.append(dict(seg_type=t, n_bins=len(g),
                         kind="FP rate" if t.startswith("ns") else "FN rate",
                         model=g.err.mean(), fill_all=(g.fill_all != g.y).mean(),
                         weight=CONTEXT_W[t]))
    strat = pd.DataFrame(rows).set_index("seg_type").reindex(list(CONTEXT_W)).dropna(how="all")
    add(fmt(strat))

    # 4. collar
    add("\n[4] ẢNH HƯỞNG CỦA LỖI BIÊN (collar)")
    nc = bins[~bins.in_collar]
    c = pd.DataFrame({"toàn bộ bin": metrics(bins.y, bins.pred),
                      f"bỏ ±{cfg['collar_bins']} bin quanh biên": metrics(nc.y, nc.pred)}).T
    add(fmt(c[["F1", "specificity", "recall"]]))
    add(f"Bin lỗi nằm trong collar: {bins[bins.err == 1].in_collar.mean():.1%} tổng bin lỗi")

    # 5. severity levels
    add("\n[5] PHÂN LOẠI MỨC ĐỘ NGHIÊM TRỌNG (theo đoạn lỗi liên tiếp)")
    if len(errs):
        lv = errs.groupby("level").agg(n_runs=("n_bins", "size"), n_bins=("n_bins", "sum"),
                                       FP=("kind", lambda s: (s == "FP").sum()),
                                       FN=("kind", lambda s: (s == "FN").sum()),
                                       confident=("confident", "sum"))
        lv["share_bins"] = lv.n_bins / lv.n_bins.sum()
        add(fmt(lv))
        severe = lv.loc[lv.index.str.match(r"L[34]"), "n_bins"].sum() / lv.n_bins.sum()
        mild = lv.loc[lv.index.str.match(r"L[01]"), "n_bins"].sum() / lv.n_bins.sum()
        add(f"→ Lỗi nhẹ (L0+L1): {mild:.1%} | Lỗi nặng (L3+L4): {severe:.1%}")
    else:
        add("Không có lỗi.")

    # 6. weighted
    raw_er = bins.err.mean(); w_er = bins.w_err.sum() / len(bins)
    add("\n[6] TỶ LỆ LỖI CÓ TRỌNG SỐ NGỮ CẢNH")
    add(f"Error rate thô: {raw_er:.3%} | Error rate có trọng số: {w_er:.3%} "
        f"(lỗi 'thực chất' ≈ {w_er / raw_er if raw_er else 0:.0%} lỗi thô)")

    # 7. per category + CI
    add("\n[7] THEO CATEGORY (CI 95% bootstrap theo file)")
    rows = []
    for cat, g in bins.groupby("category"):
        m = metrics(g.y, g.pred)
        (f_lo, f_hi), (p_lo, p_hi) = ci[cat]
        nl = g[g.seg_type == "ns_long"]; ig = g[g.seg_type == "ns_intra_gap"]
        rows.append(dict(category=cat, files=g.file.nunique(), F1=m["F1"],
                         F1_CI=f"[{f_lo:.3f},{f_hi:.3f}]", spec=m["specificity"],
                         FP_ns_long=nl.err.mean() if len(nl) else np.nan,
                         FP_ns_long_CI=f"[{p_lo:.3f},{p_hi:.3f}]",
                         FP_intra_gap=ig.err.mean() if len(ig) else np.nan,
                         w_err=g.w_err.sum() / len(g)))
        if f_hi - f_lo > 0.10:
            v.append(f"Category {cat}: CI F1 rộng {f_hi - f_lo:.3f} → chưa đủ file để kết luận.")
    add(fmt(pd.DataFrame(rows).set_index("category")))

    # 8. worst files & confident errors
    add("\n[8] TOP 10 FILE CÓ LỖI TRỌNG SỐ CAO NHẤT")
    pf = bins.groupby(["category", "file"]).agg(dur_s=("bin", lambda x: len(x) * b),
                                                err=("err", "mean"),
                                                w_err=("w_err", "mean")).sort_values("w_err", ascending=False)
    add(fmt(pf.head(10)))
    add("\n[9] LỖI TỰ TIN CAO — ưu tiên nghe lại (có thể là lỗi nhãn GT)")
    if len(errs) and errs.confident.any():
        ce = errs[errs.confident].assign(d=lambda d: (d.mean_score - 0.5).abs()) \
            .sort_values(["d", "n_bins"], ascending=False).head(15)
        add(fmt(ce[["file", "t_start", "t_end", "kind", "seg_type", "level", "mean_score"]], index=False))
    else:
        add(f"Không có lỗi vượt ngưỡng tự tin ({cfg['conf_hi']}/{cfg['conf_lo']}).")
    if len(errs) and errs.mean_score.isna().all():
        add("File model không có cột score → không xếp được theo độ tự tin.")
    elif len(errs):
        add("\nTop 10 lỗi có score xa ngưỡng nhất (nên nghe lại):")
        far = errs.assign(d=(errs.mean_score - cfg["threshold"]).abs()).sort_values("d", ascending=False).head(10)
        add(fmt(far[["file", "t_start", "t_end", "kind", "seg_type", "level", "mean_score"]], index=False))

    # verdict
    n_long = int((segs.type == "ns_long").sum())
    if n_long < cfg["min_events"]:
        v.append(f"Chỉ có {n_long} đoạn non-speech dài (<{cfg['min_events']}) → chưa đủ để đánh giá khả năng loại nhiễu.")
    if "ns_intra_gap" in st.index and st.loc["ns_intra_gap", "share_of_ns"] > 0.5:
        v.append("Hơn 50% thời lượng non-speech là khoảng ngắt ngắn → specificity bị chi phối bởi vùng dễ/ít quan trọng.")
    n_short = int((segs.type == "sp_short").sum())
    if n_short < cfg["min_events"]:
        v.append(f"Chỉ có {n_short} đoạn speech ngắn → chưa kiểm tra được khả năng bắt câu ngắn.")
    cnt = segs.pivot_table(index="category", columns="type", values="seg", aggfunc="size", fill_value=0)
    for cat, r in cnt.iterrows():
        lack = [t for t in ["ns_long", "sp_short"] if r.get(t, 0) < cfg["min_events"]]
        if lack:
            v.append(f"Category {cat}: thiếu đoạn " + ", ".join(f"{t} ({int(r.get(t, 0))})" for t in lack)
                     + " → metric của nhóm này ở category này chưa đáng tin.")
    add("\n[KẾT LUẬN TỰ ĐỘNG]")
    add("\n".join(f"- {x}" for x in v) if v else "- Không phát hiện vấn đề lớn về độ phủ dataset.")
    return "\n".join(L), v


# ----------------------------------------------------------------------------- excel
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3864")
SEC_FILL = PatternFill("solid", start_color="D9E1F2")
BAD_FILL = PatternFill("solid", start_color="F8CBAD")
BLUE = "0000FF"
EXCEL_MAX_ROWS = 1_048_576
SEG_DESC = {
    "ns_intra_gap": "Non-speech kẹp giữa 2 đoạn speech, ≤ intra_gap_max (khoảng ngắt tự nhiên)",
    "ns_edge": "Non-speech ở đầu/cuối file",
    "ns_turn_gap": "Non-speech giữa 2 đoạn speech, trong (intra_gap_max, turn_gap_max]",
    "ns_long": "Non-speech > turn_gap_max – vùng kiểm tra khả năng loại nhiễu",
    "sp_short": "Speech ≤ short_speech_max (câu ngắn: ừ, vâng...)",
    "sp_normal": "Speech dài hơn short_speech_max",
}
LEVELS = [
    ("L0_gap_fill", "FP lấp khoảng ngắt ngắn giữa speech – không nghiêm trọng"),
    ("L1_boundary", "Lệch biên trong collar – chủ yếu do lượng tử hoá 0.5s"),
    ("L2_short", "Lỗi 1–2 bin liên tiếp, xa biên"),
    ("L3_segment", "Lỗi ≥ 3 bin liên tiếp"),
    ("L4_missed_whole_speech", "Bỏ sót trọn một đoạn speech"),
    ("L4_whole_long_ns_as_speech", "Nhận trọn một đoạn non-speech dài là speech"),
]


def _py(v):
    if isinstance(v, (bool, np.bool_)):
        return int(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return None if np.isnan(v) else float(v)
    return v


class Ref:
    """Trả về range tuyệt đối của 1 cột trong sheet dữ liệu, ví dụ Bins!$E$2:$E$1741."""
    def __init__(self, sheet, df):
        self.sheet, self.df, self.n = sheet, df, max(len(df), 1) + 1

    def __getitem__(self, col):
        L = get_column_letter(self.df.columns.get_loc(col) + 1)
        return f"{self.sheet}!${L}$2:${L}${self.n}"


def _cifs(*pairs):
    return "COUNTIFS(" + ",".join(f"{r},{c}" for r, c in pairs) + ")"


def write_excel(path, bins, segs, errs, cfg, ci, verdicts):
    if len(bins) + 1 > EXCEL_MAX_ROWS:
        raise SystemExit(f"{len(bins)} bin vượt giới hạn dòng Excel; hãy chạy theo từng category.")
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONT, size=10)
    wb.calculation.fullCalcOnLoad = True
    ws = wb.active
    ws.title = "Summary"

    # ---- dữ liệu cho các sheet
    B = bins.assign(fill_err=(bins.fill_all != bins.y).astype(int),
                    in_collar=bins.in_collar.astype(int))[
        [c for c in ["category", "file", "bin", "t0", "y", "pred", "pred_raw", "fill_all", "all_speech",
                     "score", "seg", "seg_type", "dist", "err", "fill_err", "in_collar"]
         if cfg["has_raw"] or c != "pred_raw"]]
    S = segs[["category", "file", "seg", "type", "label", "b0", "b1", "dur", "pred_speech_frac"]]
    ecols = ["category", "file", "t_start", "t_end", "n_bins", "kind", "seg_type", "level",
             "mean_score", "confident"]
    E = errs[ecols].assign(confident=errs.confident.astype(int)) if len(errs) else pd.DataFrame(columns=ecols)
    F = (bins.groupby(["category", "file"]).w_err.mean().sort_values(ascending=False)
         .reset_index()[["category", "file"]])
    F = F.assign(n_bins=None, dur_s=None, err_rate=None, w_err_rate=None)
    rb, rs, re_, rf = Ref("Bins", B), Ref("Segments", S), Ref("Error_runs", E), Ref("Per_file", F)

    def put(r, c, v, fmt=None, bold=False, color=None, fill=None, sheet=ws, size=10, italic=False):
        cell = sheet.cell(r, c, v)
        cell.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        return cell

    def header(r, labels, c0=1):
        for j, t in enumerate(labels):
            c = put(r, c0 + j, t, bold=True, color="FFFFFF", fill=HDR_FILL)
            c.alignment = Alignment(wrap_text=True, vertical="center")

    def section(r, title):
        for c in range(1, 16):
            ws.cell(r, c).fill = SEC_FILL
        put(r, 1, title, bold=True, fill=SEC_FILL, size=11)

    # ---- tham số (cột Q:R, chữ xanh = nhập tay)
    P, PL = {}, {}
    put(1, 1, "BÁO CÁO PHÂN TÍCH DATASET & MỨC ĐỘ LỖI VAD", bold=True, size=14)
    put(2, 1, "Chữ xanh = giá trị nhập hoặc cố định (tính bằng Python lúc chạy); chữ đen = công thức "
              "tự tính từ các sheet Bins / Segments / Error_runs / Per_file.", italic=True)
    pr = 4
    put(pr, 17, "THAM SỐ", bold=True, fill=SEC_FILL); put(pr, 18, None, fill=SEC_FILL); pr += 1
    put(pr, 17, "Chế độ model input")
    put(pr, 18, "raw" if cfg["has_raw"] else "đã rebin", color=BLUE); pr += 1
    pipe = [("threshold", "Ngưỡng Youden"), ("hop", "Hop (s)"), ("pad_before", "Pad trước (s)"),
            ("pad_after", "Pad sau (s)"), ("merge_gap", "Merge gap (s)")] if cfg["has_raw"] else \
           [("threshold", "Ngưỡng (khi không có nhãn)")]
    params = pipe + [("bin", "Bin GT (s)"),
              ("intra_gap_max", "intra_gap_max (s)"), ("turn_gap_max", "turn_gap_max (s)"),
              ("short_speech_max", "short_speech_max (s)"), ("collar_bins", "Collar (bin)"),
              ("conf_hi", "Ngưỡng tự tin FP"), ("conf_lo", "Ngưỡng tự tin FN"),
              ("min_events", "Số đoạn tối thiểu / nhóm")]
    for k, lab in params:
        put(pr, 17, lab); put(pr, 18, cfg[k], color=BLUE)
        PL[k] = f"$R${pr}"; P[k] = f"Summary!$R${pr}"; pr += 1
    c = ws.cell(5, 17)
    c.comment = Comment("Các tham số pipeline (ngưỡng → phân loại đoạn) chỉ ghi lại cấu hình đã chạy; "
                        "đổi chúng cần chạy lại script. Riêng min_events, trọng số và boundary "
                        "factor được công thức dùng trực tiếp.", "script")
    pr += 1
    put(pr, 17, "TRỌNG SỐ NGỮ CẢNH", bold=True, fill=SEC_FILL); put(pr, 18, None, fill=SEC_FILL); pr += 1
    for t, w in CONTEXT_W.items():
        put(pr, 17, t); put(pr, 18, w, color=BLUE)
        PL["w_" + t] = f"$R${pr}"; P["w_" + t] = f"Summary!$R${pr}"; pr += 1
    put(pr, 17, "Hệ số lỗi biên (collar)"); put(pr, 18, BOUNDARY_FACTOR, color=BLUE)
    PL["bf"] = f"$R${pr}"; P["bf"] = f"Summary!$R${pr}"

    def wexpr(extra=""):
        terms = []
        for t in CONTEXT_W:
            base = f'{rb["seg_type"]},"{t}",{rb["err"]},1'
            terms.append(f'{P["w_" + t]}*(COUNTIFS({base},{rb["in_collar"]},0{extra})'
                         f'+{P["bf"]}*COUNTIFS({base},{rb["in_collar"]},1{extra}))')
        return "+".join(terms)

    def metric_row(r, label, predcol, extra=()):
        put(r, 1, label)
        put(r, 2, "=" + _cifs((rb["y"], 1), (rb[predcol], 1), *extra), "#,##0")
        put(r, 3, "=" + _cifs((rb["y"], 0), (rb[predcol], 1), *extra), "#,##0")
        put(r, 4, "=" + _cifs((rb["y"], 1), (rb[predcol], 0), *extra), "#,##0")
        put(r, 5, "=" + _cifs((rb["y"], 0), (rb[predcol], 0), *extra), "#,##0")
        put(r, 6, f'=IFERROR(B{r}/(B{r}+C{r}),"")', "0.000")
        put(r, 7, f'=IFERROR(B{r}/(B{r}+D{r}),"")', "0.000")
        put(r, 8, f'=IFERROR(2*B{r}/(2*B{r}+C{r}+D{r}),"")', "0.000")
        put(r, 9, f'=IFERROR(E{r}/(E{r}+C{r}),"")', "0.000")
        put(r, 10, f'=IFERROR((B{r}+E{r})/(B{r}+C{r}+D{r}+E{r}),"")', "0.000")

    MHDR = ["TP", "FP", "FN", "TN", "Precision", "Recall", "F1", "Specificity", "Accuracy"]
    r = 4
    # ---- tổng quan
    section(r, "TỔNG QUAN"); r += 1
    for lab, f, fmt in [("Số file", f"=COUNTA({rf['file']})", "0"),
                        ("Số category", f"=SUMPRODUCT(1/COUNTIF({rf['category']},{rf['category']}))", "0"),
                        ("Số bin", f"=COUNT({rb['y']})", "#,##0"),
                        ("Tổng thời lượng (phút)", f"=COUNT({rb['y']})*{PL['bin']}/60", "0.0"),
                        ("Tỷ lệ speech", f"=AVERAGE({rb['y']})", "0.0%")]:
        put(r, 1, lab); put(r, 2, f, fmt); r += 1

    # ---- [1] cấu trúc
    r += 1; section(r, "[1] CẤU TRÚC ĐOẠN TRONG GROUND TRUTH"); r += 1
    header(r, ["Loại đoạn", "Số đoạn", "Tổng (s)", "TB (s)", "% thời lượng NS", "Đủ mẫu?", "Mô tả"]); r += 1
    r0 = r
    for t in CONTEXT_W:
        put(r, 1, t)
        put(r, 2, f'=COUNTIFS({rs["type"]},"{t}")', "0")
        put(r, 3, f'=SUMIFS({rs["dur"]},{rs["type"]},"{t}")', "0.0")
        put(r, 4, f'=IFERROR(C{r}/B{r},"")', "0.00")
        if t.startswith("ns"):
            put(r, 5, f'=IFERROR(C{r}/SUMIFS({rs["dur"]},{rs["label"]},0),"")', "0.0%")
        put(r, 6, f'=IF(B{r}>={PL["min_events"]},"OK","THIẾU")')
        put(r, 7, SEG_DESC[t], italic=True)
        r += 1
    ws.conditional_formatting.add(f"F{r0}:F{r - 1}", CellIsRule(operator="equal", formula=['"THIẾU"'],
                                                               fill=BAD_FILL))
    r += 1
    put(r, 1, "Số đoạn theo category (ô đỏ = dưới ngưỡng tối thiểu)", bold=True); r += 1
    types = list(CONTEXT_W)
    header(r, ["Category"] + types); r += 1
    r0 = r
    for cat in sorted(bins.category.unique()):
        put(r, 1, cat)
        for j, t in enumerate(types):
            put(r, 2 + j, f'=COUNTIFS({rs["category"]},"{cat}",{rs["type"]},"{t}")', "0")
        r += 1
    for t in ["ns_long", "sp_short"]:
        L = get_column_letter(2 + types.index(t))
        ws.conditional_formatting.add(f"{L}{r0}:{L}{r - 1}", CellIsRule(
            operator="lessThan", formula=[PL["min_events"]], fill=BAD_FILL))

    # ---- [2] baseline
    r += 1; section(r, "[2] SO VỚI BASELINE NGÂY THƠ"); r += 1
    header(r, ["Predictor"] + MHDR); r += 1
    rows = {}
    preds = [("pred", "Model (post-proc)" if cfg["has_raw"] else "Model (đã rebin)"),
             ("pred_raw", "Model raw (không post-proc)"),
             ("fill_all", "Fill-all (onset đầu → offset cuối)"), ("all_speech", "All-speech")]
    for col, lab in preds:
        if col == "pred_raw" and not cfg["has_raw"]:
            continue
        metric_row(r, lab, col); rows[col] = r; r += 1
    put(r, 1, "Chênh F1: model − fill-all", bold=True)
    put(r, 2, f"=H{rows['pred']}-H{rows['fill_all']}", "+0.000;-0.000")
    put(r, 3, f'=IF(B{r}<0.05,"Dataset DỄ: baseline gần bằng model","Model vượt baseline rõ rệt")')
    r += 1

    # ---- [3] phân tầng
    r += 1; section(r, "[3] TỶ LỆ LỖI PHÂN TẦNG THEO NGỮ CẢNH"); r += 1
    header(r, ["Loại đoạn", "Loại lỗi", "Số bin", "Lỗi – model", "Lỗi – fill-all", "Trọng số"]); r += 1
    for t in CONTEXT_W:
        put(r, 1, t); put(r, 2, "FP rate" if t.startswith("ns") else "FN rate")
        put(r, 3, f'=COUNTIFS({rb["seg_type"]},"{t}")', "#,##0")
        put(r, 4, f'=IFERROR(COUNTIFS({rb["seg_type"]},"{t}",{rb["err"]},1)/C{r},"")', "0.0%")
        put(r, 5, f'=IFERROR(COUNTIFS({rb["seg_type"]},"{t}",{rb["fill_err"]},1)/C{r},"")', "0.0%")
        put(r, 6, f"={PL['w_' + t]}", "0.00")
        r += 1

    # ---- [4] collar
    r += 1; section(r, "[4] ẢNH HƯỞNG CỦA LỖI BIÊN (COLLAR)"); r += 1
    header(r, ["Phạm vi"] + MHDR); r += 1
    metric_row(r, "Toàn bộ bin", "pred"); r += 1
    metric_row(r, "Bỏ các bin trong collar", "pred", extra=[(rb["in_collar"], 0)]); r += 1
    put(r, 1, "% bin lỗi nằm trong collar", bold=True)
    put(r, 2, f'=IFERROR(COUNTIFS({rb["err"]},1,{rb["in_collar"]},1)/COUNTIFS({rb["err"]},1),"")', "0.0%")
    r += 1

    # ---- [5] mức độ
    r += 1; section(r, "[5] PHÂN LOẠI MỨC ĐỘ NGHIÊM TRỌNG"); r += 1
    header(r, ["Mức", "Số đoạn lỗi", "Số bin", "FP", "FN", "Lỗi tự tin", "% bin lỗi", "Mô tả"]); r += 1
    r0 = r
    for lv, desc in LEVELS:
        put(r, 1, lv)
        put(r, 2, f'=COUNTIFS({re_["level"]},"{lv}")', "0")
        put(r, 3, f'=SUMIFS({re_["n_bins"]},{re_["level"]},"{lv}")', "0")
        put(r, 4, f'=COUNTIFS({re_["level"]},"{lv}",{re_["kind"]},"FP")', "0")
        put(r, 5, f'=COUNTIFS({re_["level"]},"{lv}",{re_["kind"]},"FN")', "0")
        put(r, 6, f'=COUNTIFS({re_["level"]},"{lv}",{re_["confident"]},1)', "0")
        put(r, 7, f'=IFERROR(C{r}/SUM($C${r0}:$C${r0 + len(LEVELS) - 1}),"")', "0.0%")
        put(r, 8, desc, italic=True)
        r += 1
    put(r, 1, "Lỗi nhẹ (L0 + L1)", bold=True); put(r, 2, f"=G{r0}+G{r0 + 1}", "0.0%"); r += 1
    put(r, 1, "Lỗi nặng (L3 + L4)", bold=True)
    put(r, 2, f"=G{r0 + 3}+G{r0 + 4}+G{r0 + 5}", "0.0%"); r += 1

    # ---- [6] trọng số
    r += 1; section(r, "[6] TỶ LỆ LỖI CÓ TRỌNG SỐ NGỮ CẢNH (đổi trọng số ở cột R)"); r += 1
    header(r, ["Loại đoạn", "Trọng số", "Bin lỗi ngoài collar", "Bin lỗi trong collar",
               "Đóng góp có trọng số"]); r += 1
    r0 = r
    for t in CONTEXT_W:
        put(r, 1, t); put(r, 2, f"={PL['w_' + t]}", "0.00")
        put(r, 3, f'=COUNTIFS({rb["seg_type"]},"{t}",{rb["err"]},1,{rb["in_collar"]},0)', "#,##0")
        put(r, 4, f'=COUNTIFS({rb["seg_type"]},"{t}",{rb["err"]},1,{rb["in_collar"]},1)', "#,##0")
        put(r, 5, f"=B{r}*(C{r}+{PL['bf']}*D{r})", "#,##0.0")
        r += 1
    put(r, 1, "Error rate thô", bold=True)
    put(r, 2, f'=IFERROR(COUNTIFS({rb["err"]},1)/COUNT({rb["y"]}),"")', "0.00%"); raw_r = r; r += 1
    put(r, 1, "Error rate có trọng số", bold=True)
    put(r, 2, f'=IFERROR(SUM(E{r0}:E{r - 2})/COUNT({rb["y"]}),"")', "0.00%"); w_r = r; r += 1
    put(r, 1, "Lỗi thực chất / lỗi thô", bold=True)
    put(r, 2, f'=IFERROR(B{w_r}/B{raw_r},"")', "0%"); r += 1

    # ---- [7] theo category
    r += 1; section(r, "[7] THEO CATEGORY"); r += 1
    header(r, ["Category", "Số file", "TP", "FP", "FN", "TN", "F1", "F1 CI 2.5%", "F1 CI 97.5%",
               "Specificity", "FP ns_long", "ns_long CI 2.5%", "ns_long CI 97.5%", "FP intra_gap",
               "Error có trọng số"]); r += 1
    note = (f"Khoảng tin cậy 95% bootstrap theo file ({cfg['n_boot']} lần), tính bằng Python lúc chạy "
            "script – không tự cập nhật khi sửa dữ liệu.")
    ws.cell(r - 1, 8).comment = Comment(note, "script")
    ws.cell(r - 1, 12).comment = Comment(note, "script")
    for cat in sorted(bins.category.unique()):
        cc = (rb["category"], f'"{cat}"')
        put(r, 1, cat)
        put(r, 2, f'=COUNTIFS({rf["category"]},"{cat}")', "0")
        put(r, 3, "=" + _cifs(cc, (rb["y"], 1), (rb["pred"], 1)), "#,##0")
        put(r, 4, "=" + _cifs(cc, (rb["y"], 0), (rb["pred"], 1)), "#,##0")
        put(r, 5, "=" + _cifs(cc, (rb["y"], 1), (rb["pred"], 0)), "#,##0")
        put(r, 6, "=" + _cifs(cc, (rb["y"], 0), (rb["pred"], 0)), "#,##0")
        put(r, 7, f'=IFERROR(2*C{r}/(2*C{r}+D{r}+E{r}),"")', "0.000")
        (f_lo, f_hi), (p_lo, p_hi) = ci[cat]
        put(r, 8, _py(f_lo), "0.000", color=BLUE); put(r, 9, _py(f_hi), "0.000", color=BLUE)
        put(r, 10, f'=IFERROR(F{r}/(F{r}+D{r}),"")', "0.000")
        for col, t in [(11, "ns_long"), (14, "ns_intra_gap")]:
            put(r, col, f'=IFERROR({_cifs(cc, (rb["seg_type"], chr(34) + t + chr(34)), (rb["err"], 1))}'
                        f'/{_cifs(cc, (rb["seg_type"], chr(34) + t + chr(34)))},"")', "0.0%")
        put(r, 12, _py(p_lo), "0.0%", color=BLUE); put(r, 13, _py(p_hi), "0.0%", color=BLUE)
        put(r, 15, f'=IFERROR(({wexpr("," + rb["category"] + "," + chr(34) + cat + chr(34))})'
                   f'/COUNTIFS({rb["category"]},"{cat}"),"")', "0.00%")
        r += 1

    # ---- [8] top file
    r += 1; section(r, "[8] TOP 10 FILE CÓ LỖI TRỌNG SỐ CAO NHẤT (thứ tự lúc chạy script)"); r += 1
    header(r, ["Category", "File", "Thời lượng (s)", "Error rate", "Error có trọng số"]); r += 1
    for i in range(min(10, len(F))):
        pr_ = i + 2
        put(r, 1, f"=Per_file!A{pr_}"); put(r, 2, f"=Per_file!B{pr_}")
        put(r, 3, f"=Per_file!D{pr_}", "0.0"); put(r, 4, f"=Per_file!E{pr_}", "0.0%")
        put(r, 5, f"=Per_file!F{pr_}", "0.0%")
        r += 1

    # ---- [9] nghe lại
    r += 1; section(r, "[9] DANH SÁCH NÊN NGHE LẠI – lỗi có score xa ngưỡng nhất (có thể là lỗi nhãn GT)"); r += 1
    header(r, ["File", "Bắt đầu (s)", "Kết thúc (s)", "Loại lỗi", "Loại đoạn", "Mức", "Score TB",
               "Tự tin?"]); r += 1
    if len(errs) and errs.mean_score.isna().all():
        put(r, 1, "File model không có cột score → không xếp được theo độ tự tin.", italic=True); r += 1
    elif len(errs):
        far = errs.assign(d=(errs.mean_score - cfg["threshold"]).abs()) \
            .sort_values(["confident", "d"], ascending=False).head(15)
        for row in far.itertuples():
            for j, v in enumerate([row.file, row.t_start, row.t_end, row.kind, row.seg_type,
                                   row.level, row.mean_score, "Có" if row.confident else ""]):
                put(r, 1 + j, _py(v), "0.00" if j in (1, 2) else "0.000" if j == 6 else None, color=BLUE)
            r += 1

    # ---- [10] kết luận
    r += 1; section(r, "[10] KẾT LUẬN TỰ ĐỘNG (sinh bởi Python lúc chạy script)"); r += 1
    for t in verdicts or ["Không phát hiện vấn đề lớn về độ phủ dataset."]:
        put(r, 1, "• " + t); r += 1

    ws.column_dimensions["A"].width = 36
    for c in range(2, 16):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.column_dimensions["Q"].width = 26
    ws.column_dimensions["R"].width = 10
    ws.freeze_panes = "B4"

    # ---- sheet dữ liệu
    def data_sheet(name, df):
        sh = wb.create_sheet(name)
        sh.append(list(df.columns))
        for row in df.itertuples(index=False):
            sh.append([_py(v) for v in row])
        for c in sh[1]:
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF"); c.fill = HDR_FILL
        sh.freeze_panes = "A2"
        sh.auto_filter.ref = sh.dimensions
        for i, col in enumerate(df.columns, 1):
            sh.column_dimensions[get_column_letter(i)].width = max(10, len(col) + 4)
        return sh

    pf = data_sheet("Per_file", F)
    for i in range(len(F)):
        rr = i + 2
        pf.cell(rr, 3, f'=COUNTIFS({rb["file"]},B{rr})').number_format = "0"
        pf.cell(rr, 4, f"=C{rr}*{P['bin']}").number_format = "0.0"
        pf.cell(rr, 5, f'=IFERROR(COUNTIFS({rb["file"]},B{rr},{rb["err"]},1)/C{rr},"")').number_format = "0.0%"
        pf.cell(rr, 6, f'=IFERROR(({wexpr("," + rb["file"] + ",B" + str(rr))})/C{rr},"")').number_format = "0.0%"
    pf.column_dimensions["B"].width = 18
    data_sheet("Error_runs", E)
    data_sheet("Segments", S)
    data_sheet("Bins", B)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--model_root", default="Model")
    ap.add_argument("--gt_root", default="Groundtruth")
    ap.add_argument("--out", default="vad_report.xlsx")
    ap.add_argument("--threshold", type=float, default=CFG["threshold"])
    ap.add_argument("--model_mode", choices=["auto", "raw", "rebinned"], default="auto",
                    help="raw = output YAMNet chưa xử lý; rebinned = đã rebin 0.5s")
    a = ap.parse_args()
    cfg = dict(CFG, threshold=a.threshold)
    out = Path(a.out)
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    pairs = match_files(a.model_root, a.gt_root)
    if not pairs:
        raise SystemExit("Không ghép được cặp file Model/GT nào.")
    mode = a.model_mode
    if mode == "auto":
        mode = "rebinned" if is_frame_grid(read_frames(pairs[0][2]), cfg["bin"]) else "raw"
    cfg.update(mode=mode, has_raw=(mode == "raw"))
    print(f"Chế độ model input: {mode}")
    B, S, E = [], [], []
    for cat, name, mf, gf in pairs:
        b_, s_, e_ = analyze_file(cat, name, mf, gf, cfg)
        B.append(b_); S.append(s_); E.append(e_)
    bins, segs = pd.concat(B, ignore_index=True), pd.concat(S, ignore_index=True)
    errs = pd.concat([e for e in E if len(e)], ignore_index=True) if any(len(e) for e in E) else pd.DataFrame()

    ci = {cat: boot_ci(g, cfg) for cat, g in bins.groupby("category")}
    report, verdicts = build_report(bins, segs, errs, cfg, ci)
    print(report)
    write_excel(out, bins, segs, errs, cfg, ci, verdicts)
    print(f"\nĐã lưu: {out}  (sheet Summary, Per_file, Error_runs, Segments, Bins)")


if __name__ == "__main__":
    main()
